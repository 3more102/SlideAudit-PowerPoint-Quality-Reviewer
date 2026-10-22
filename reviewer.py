"""Explainable, non-AI PowerPoint presentation quality checks.

All scores are heuristics. A PPTX master/theme can hide effective typography or
color; we deliberately leave such properties unassessed rather than guess.
"""
from __future__ import annotations

from dataclasses import asdict, dataclass
from io import BytesIO
from statistics import median
from typing import Any
from zipfile import BadZipFile, ZipFile
import json

from pptx import Presentation
from pptx.enum.dml import MSO_COLOR_TYPE, MSO_FILL_TYPE
from pptx.enum.shapes import MSO_SHAPE_TYPE

WEIGHTS = {
    "Typography": 20,
    "Content density": 20,
    "Layout": 20,
    "Visual hierarchy": 15,
    "Consistency": 15,
    "Contrast": 10,
}
MAX_FILE_BYTES = 50 * 1024 * 1024
MAX_UNCOMPRESSED_BYTES = 180 * 1024 * 1024
MAX_MEMBERS = 7000


@dataclass
class Finding:
    category: str
    severity: str
    impact: int
    headline: str
    evidence: str
    fix: str
    shape: str = ""


@dataclass
class SlideReport:
    number: int
    title: str
    score: int
    metrics: dict[str, Any]
    category_scores: dict[str, int]
    findings: list[Finding]


@dataclass
class DeckReport:
    name: str
    score: int
    slide_count: int
    slides: list[SlideReport]
    category_scores: dict[str, int]
    executive_notes: list[str]
    limitations: list[str]
    rubric: dict[str, int]

    def to_dict(self) -> dict[str, Any]:
        return asdict(self)

    def to_json(self) -> str:
        return json.dumps(self.to_dict(), indent=2, ensure_ascii=False)

    def to_markdown(self) -> str:
        out = [f"# SlideAudit report: {self.name}", "", f"Overall heuristic score: **{self.score}/100**", f"Slides evaluated: {self.slide_count}", "", "## Highest-priority edits"]
        out.extend(f"- {note}" for note in self.executive_notes)
        out += ["", "## Category scores"]
        out.extend(f"- {k}: {self.category_scores[k]}/{v}" for k, v in self.rubric.items())
        for slide in self.slides:
            out += ["", f"## Slide {slide.number}: {slide.title} — {slide.score}/100"]
            if not slide.findings:
                out.append("No issues detected by the supported automated checks.")
            for f in slide.findings:
                out += [f"- **[{f.severity}] {f.headline}** ({f.category}, −{f.impact} possible points)", f"  - Evidence: {f.evidence}", f"  - Edit: {f.fix}"]
        out += ["", "## Limits of the automated assessment"]
        out.extend(f"- {x}" for x in self.limitations)
        return "\n".join(out) + "\n"


def _validate_pptx(data: bytes) -> None:
    if not data:
        raise ValueError("The uploaded file is empty.")
    if len(data) > MAX_FILE_BYTES:
        raise ValueError("The file exceeds the 50 MB safety limit.")
    try:
        with ZipFile(BytesIO(data)) as zip_file:
            files = zip_file.infolist()
            if len(files) > MAX_MEMBERS or sum(f.file_size for f in files) > MAX_UNCOMPRESSED_BYTES:
                raise ValueError("The PPTX expands beyond its safety limits.")
            if "[Content_Types].xml" not in zip_file.namelist() or "ppt/presentation.xml" not in zip_file.namelist():
                raise ValueError("This file is not a valid .pptx presentation.")
            if any(f.flag_bits & 0x1 for f in files):
                raise ValueError("Encrypted PowerPoint files are not supported.")
    except BadZipFile as exc:
        raise ValueError("The file is not a valid .pptx archive. Export it as .pptx first.") from exc


def _shape_name(shape: Any) -> str:
    return getattr(shape, "name", "Unnamed object")


def _shape_text(shape: Any) -> str:
    if getattr(shape, "has_text_frame", False):
        return (shape.text or "").strip()
    if getattr(shape, "has_table", False):
        return "\n".join(cell.text for row in shape.table.rows for cell in row.cells).strip()
    return ""


def _paragraphs(shape: Any) -> list[Any]:
    if getattr(shape, "has_text_frame", False):
        return list(shape.text_frame.paragraphs)
    if getattr(shape, "has_table", False):
        return [p for row in shape.table.rows for cell in row.cells for p in cell.text_frame.paragraphs]
    return []


def _explicit_fonts(shape: Any) -> list[float]:
    fonts = []
    for p in _paragraphs(shape):
        for run in p.runs:
            if run.text.strip() and run.font.size is not None:
                fonts.append(round(run.font.size.pt, 1))
        if p.text.strip() and not p.runs and p.font.size is not None:
            fonts.append(round(p.font.size.pt, 1))
    return fonts


def _font_families(shape: Any) -> list[str]:
    names = []
    for p in _paragraphs(shape):
        for run in p.runs:
            if run.text.strip() and run.font.name:
                names.append(run.font.name.strip())
        if p.text.strip() and not p.runs and p.font.name:
            names.append(p.font.name.strip())
    return names


def _rgb_from_color(color: Any) -> tuple[int, int, int] | None:
    try:
        if color.type == MSO_COLOR_TYPE.RGB:
            value = color.rgb
            return value[0], value[1], value[2]
    except (AttributeError, TypeError, ValueError):
        pass
    return None


def _solid_fill(shape_or_slide: Any) -> tuple[int, int, int] | None:
    try:
        fill = shape_or_slide.fill if hasattr(shape_or_slide, "fill") else shape_or_slide.background.fill
        if fill.type == MSO_FILL_TYPE.SOLID:
            return _rgb_from_color(fill.fore_color)
    except (AttributeError, TypeError, ValueError):
        pass
    return None


def _luminance(rgb: tuple[int, int, int]) -> float:
    linear = []
    for v in rgb:
        v = v / 255.0
        linear.append(v / 12.92 if v <= 0.04045 else ((v + 0.055) / 1.055) ** 2.4)
    return 0.2126 * linear[0] + 0.7152 * linear[1] + 0.0722 * linear[2]


def contrast_ratio(a: tuple[int, int, int], b: tuple[int, int, int]) -> float:
    bright, dark = sorted((_luminance(a), _luminance(b)), reverse=True)
    return round((bright + 0.05) / (dark + 0.05), 2)


def _contains(outer: Any, inner: Any) -> bool:
    return (outer.left <= inner.left and outer.top <= inner.top and
            outer.left + outer.width >= inner.left + inner.width and
            outer.top + outer.height >= inner.top + inner.height)


def _behind_fill(slide: Any, shape: Any) -> tuple[int, int, int] | None:
    current = _solid_fill(shape)
    if current is not None:
        return current
    candidates = []
    for other in slide.shapes:
        if other is shape:
            break  # PowerPoint stacking order: later objects appear on top
        fill = _solid_fill(other)
        if fill is not None and _contains(other, shape):
            candidates.append((other.width * other.height, fill))
    if candidates:
        return min(candidates, key=lambda v: v[0])[1]
    return _solid_fill(slide)


def _title_candidate(slide: Any, width: int, height: int) -> Any | None:
    title = slide.shapes.title
    if title is not None and _shape_text(title):
        return title
    choices = []
    for shape in slide.shapes:
        text = _shape_text(shape)
        if not text or len(text.split()) > 24 or shape.top > height * 0.31:
            continue
        fonts = _explicit_fonts(shape)
        # Do not assume paragraph sizes inherited from master are known.
        size = max(fonts) if fonts else 0.0
        choices.append((size, -shape.top, shape))
    return max(choices, key=lambda x: (x[0], x[1]))[2] if choices else None


def _intersection_fraction(a: Any, b: Any) -> float:
    x = max(0, min(a.left + a.width, b.left + b.width) - max(a.left, b.left))
    y = max(0, min(a.top + a.height, b.top + b.height) - max(a.top, b.top))
    return (x * y) / max(1, min(a.width * a.height, b.width * b.height))


def _issue(slide_findings: list[Finding], category: str, severity: str, impact: int,
           headline: str, evidence: str, fix: str, shape: str = "") -> None:
    slide_findings.append(Finding(category, severity, impact, headline, evidence, fix, shape))


def _slide_metrics(slide: Any, width: int, height: int) -> tuple[dict[str, Any], Any | None, list[Finding]]:
    shapes = list(slide.shapes)
    texts = [(s, _shape_text(s)) for s in shapes if _shape_text(s)]
    words = sum(len(text.split()) for _, text in texts)
    lines = sum(max(1, len(text.splitlines())) for _, text in texts)
    visuals = sum(bool(s.shape_type == MSO_SHAPE_TYPE.PICTURE or getattr(s, "has_chart", False) or
                       getattr(s, "has_table", False) or s.shape_type in (MSO_SHAPE_TYPE.GROUP, MSO_SHAPE_TYPE.DIAGRAM))
                  for s in shapes)
    # A cluster of vector shapes can itself be an intentional visual diagram.
    vector_objects = sum(s.shape_type in (MSO_SHAPE_TYPE.AUTO_SHAPE, MSO_SHAPE_TYPE.FREEFORM)
                         for s in shapes)
    if vector_objects >= 2:
        visuals += 1
    explicit_sizes = [f for s, _ in texts for f in _explicit_fonts(s)]
    fonts = [n for s, _ in texts for n in _font_families(s)]
    title = _title_candidate(slide, width, height)
    findings: list[Finding] = []
    if title is None and words:
        _issue(findings, "Visual hierarchy", "High", 8, "No identifiable slide title",
               "No title placeholder or short upper-slide heading was found.",
               "Add a short, descriptive heading in the top region and apply a title layout.")
    if words > 90:
        _issue(findings, "Content density", "High", 13, "Very dense slide",
               f"{words} words (heuristic guide: ≤60).", "Split this slide into 2–3 slides or replace paragraphs with a labeled diagram.")
    elif words > 60:
        _issue(findings, "Content density", "Medium", 8, "Too much slide text",
               f"{words} words (heuristic guide: ≤60).", "Shorten paragraphs into one main takeaway and supporting bullets.")
    elif words > 45:
        _issue(findings, "Content density", "Low", 3, "Consider trimming slide text",
               f"{words} words.", "Cut non-essential words; move explanations into speaker notes.")
    if lines > 15:
        _issue(findings, "Content density", "Medium", 7, "Too many lines of text",
               f"{lines} non-empty or explicit lines.", "Increase whitespace by condensing or splitting content.")
    elif lines > 10:
        _issue(findings, "Content density", "Low", 3, "Limited whitespace between lines",
               f"{lines} non-empty or explicit lines.", "Use fewer lines and larger line spacing.")
    small = [f for f in explicit_sizes if f < 14]
    compact = [f for f in explicit_sizes if 14 <= f < 18]
    if small:
        _issue(findings, "Typography", "High", 12, "Text under 14 pt",
               f"{len(small)} explicitly sized run(s), smallest {min(small):g} pt.",
               "Increase primary slide text toward 18–24 pt where practical; reserve very small type for brief footnotes.")
    elif compact and len(compact) >= 2:
        _issue(findings, "Typography", "Medium", 6, "Repeated text under 18 pt",
               f"{len(compact)} explicitly sized runs between 14 and 18 pt.",
               "Increase body text and reduce text volume; treat dense data tables separately.")
    if title is not None:
        title_fonts = _explicit_fonts(title)
        body_fonts = [v for s, _ in texts if s is not title for v in _explicit_fonts(s)]
        if title_fonts and body_fonts and max(title_fonts) <= median(body_fonts):
            _issue(findings, "Visual hierarchy", "Medium", 6, "Title does not stand out",
                   f"Title {max(title_fonts):g} pt; median body {median(body_fonts):g} pt.",
                   "Increase title size/weight or reduce body text size to establish clear hierarchy.", _shape_name(title))
    for shape in shapes:
        if shape.left < 0 or shape.top < 0 or shape.left + shape.width > width or shape.top + shape.height > height:
            if _shape_text(shape) or shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
                _issue(findings, "Layout", "High", 12, "Content extends outside slide",
                       f"Object '{_shape_name(shape)}' crosses a slide boundary.",
                       "Move or resize the object to fit entirely within the slide canvas.", _shape_name(shape))
                if sum(f.headline == "Content extends outside slide" for f in findings) >= 2:
                    break
    text_shapes = [s for s, t in texts if t and s.width > 0 and s.height > 0]
    overlap_count = 0
    for i, a in enumerate(text_shapes):
        for b in text_shapes[i + 1:]:
            if _contains(a, b) or _contains(b, a):
                continue
            if _intersection_fraction(a, b) > 0.18:
                overlap_count += 1
    if overlap_count:
        _issue(findings, "Layout", "Medium", min(12, 6 + 2 * max(0, overlap_count - 1)),
               "Text boxes overlap", f"{overlap_count} potentially conflicting text-box pair(s).",
               "Check those boxes visually, then separate, resize, or bring the intended object forward.")
    if words >= 45 and visuals == 0:
        _issue(findings, "Visual hierarchy", "Low", 4, "Text-heavy slide without a visual",
               f"{words} words and no detected image, chart, table or grouped diagram.",
               "Consider a purposeful diagram, chart, or a single highlighted takeaway; do not add decoration unnecessarily.")
    if len(set(n.lower() for n in fonts)) > 3:
        _issue(findings, "Consistency", "Medium", 6, "Many typefaces on one slide",
               f"{len(set(n.lower() for n in fonts))} explicit font families.",
               "Use one heading family and one body family, except for meaningful code/math needs.")
    checked = 0
    min_ratio = 99.0
    for shape, text in texts:
        background = _behind_fill(slide, shape)
        if background is None:
            continue
        for p in _paragraphs(shape):
            for run in p.runs:
                if not run.text.strip():
                    continue
                try:
                    foreground = _rgb_from_color(run.font.color)
                except (AttributeError, TypeError, ValueError):
                    foreground = None
                if foreground is None:
                    continue
                checked += 1
                ratio = contrast_ratio(foreground, background)
                size = run.font.size.pt if run.font.size is not None else None
                threshold = 3.0 if size is not None and size >= 24 else 4.5
                if ratio < threshold:
                    min_ratio = min(min_ratio, ratio)
    if min_ratio < 99:
        _issue(findings, "Contrast", "High", 9, "Low contrast (explicit colors)",
               f"At least one evaluated text run had contrast {min_ratio:.2f}:1; normal text target ≥4.5:1, large text ≥3:1.",
               "Change text or solid background color to meet contrast targets; verify against actual slide visuals.")
    metrics = {"word_count": words, "line_count": lines, "visual_objects": visuals,
               "explicit_font_runs": len(explicit_sizes), "min_explicit_font_pt": min(explicit_sizes) if explicit_sizes else None,
               "contrast_runs_checked": checked, "text_objects": len(texts), "font_families": sorted(set(fonts))}
    return metrics, title, findings


def _score(findings: list[Finding]) -> tuple[int, dict[str, int]]:
    scores = {cat: max(0, weight - min(weight, sum(f.impact for f in findings if f.category == cat))) for cat, weight in WEIGHTS.items()}
    return sum(scores.values()), scores


def review_pptx(data: bytes, filename: str = "presentation.pptx") -> DeckReport:
    _validate_pptx(data)
    try:
        prs = Presentation(BytesIO(data))
    except Exception as exc:
        raise ValueError("PowerPoint could not read this presentation. Try re-exporting a .pptx.") from exc
    if not prs.slides:
        raise ValueError("This PowerPoint has no slides to evaluate.")
    if len(prs.slides) > 350:
        raise ValueError("Presentations over 350 slides are not supported.")
    width, height = prs.slide_width, prs.slide_height
    data_by_slide = []
    for slide in prs.slides:
        metrics, title, findings = _slide_metrics(slide, width, height)
        data_by_slide.append((metrics, title, findings, slide))
    # Deck-relative consistency: only inspect explicit font sizes, never infer master defaults.
    title_size_values = []
    for metrics, title, _, _ in data_by_slide:
        vals = _explicit_fonts(title) if title is not None else []
        if vals:
            title_size_values.append(max(vals))
    normal_title = median(title_size_values) if len(title_size_values) >= 3 else None
    reports = []
    for i, (metrics, title, findings, slide) in enumerate(data_by_slide, 1):
        if normal_title is not None and title is not None:
            sizes = _explicit_fonts(title)
            if sizes and abs(max(sizes) - normal_title) >= max(8, normal_title * 0.30):
                _issue(findings, "Consistency", "Medium", 6, "Title size varies across deck",
                       f"This slide {max(sizes):g} pt; deck median {normal_title:g} pt.",
                       "Apply one consistent title style across similar slide types.", _shape_name(title))
        score, per_category = _score(findings)
        reports.append(SlideReport(i, (_shape_text(title).splitlines()[0][:100] if title is not None else f"Slide {i}"),
                                   score, metrics, per_category, findings))
    score = round(sum(slide.score for slide in reports) / len(reports))
    overall = {cat: round(sum(s.category_scores[cat] for s in reports) / len(reports)) for cat in WEIGHTS}
    ranked = sorted(((f.impact, s.number, f) for s in reports for f in s.findings),
                    key=lambda t: (-t[0], t[1], t[2].headline))
    notes = [f"Slide {n}: {f.headline}. {f.fix}" for _, n, f in ranked[:7]]
    if not notes:
        notes = ["No issues flagged by these automated checks. Perform a human review for factual accuracy, visual balance, audience fit, and clarity."]
    limits = [
        "The score is a transparent rule-based heuristic, NOT a validated objective quality or academic grade.",
        "Effective theme-inherited fonts/colors and complex gradients/images are not reliably assessed; unknown values are not guessed.",
        "Contrast is tested only when both text and background have explicit RGB solid colors; visual verification is still required.",
        "Language clarity, factual accuracy, storytelling, semantic quality, animation, speaker delivery, charts and visual aesthetics need human review.",
        "Overlapping text and word-count warnings can be legitimate for intentionally designed slides; inspect the actual rendered slide.",
        "This tool reads slides locally/in the deployment process; do not upload confidential decks to third-party hosting without approval.",
    ]
    return DeckReport(filename, score, len(reports), reports, overall, notes, limits, WEIGHTS.copy())


def add_review_notes(data: bytes, report: DeckReport) -> bytes:
    """Produce a COPY of the PPTX with editable findings in each slide's speaker notes.

    Existing notes text is preserved; advanced features may not round-trip
    losslessly through python-pptx, so retain the untouched original file.
    """
    _validate_pptx(data)
    prs = Presentation(BytesIO(data))
    if len(prs.slides) != len(report.slides):
        raise ValueError("Review and slide deck must have the same number of slides.")
    for original, review in zip(prs.slides, report.slides):
        if not review.findings:
            continue
        tf = original.notes_slide.notes_text_frame
        if tf is None:
            continue
        paragraph = tf.add_paragraph()
        paragraph.text = f"[SlideAudit] Slide {review.number} — heuristic quality {review.score}/100"
        for f in review.findings:
            paragraph = tf.add_paragraph()
            paragraph.text = f"[{f.severity}] {f.headline} | Evidence: {f.evidence} | Suggested edit: {f.fix}"
    b = BytesIO()
    prs.save(b)
    return b.getvalue()
