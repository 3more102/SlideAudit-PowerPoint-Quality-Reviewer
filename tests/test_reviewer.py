from io import BytesIO
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Inches, Pt
import pytest

from reviewer import review_pptx, add_review_notes, contrast_ratio, WEIGHTS


def make_deck(*, bad=False, low_contrast=False, slides=1):
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    for idx in range(slides):
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        heading = slide.shapes.add_textbox(Inches(0.7), Inches(0.45), Inches(11), Inches(0.7))
        heading.text_frame.paragraphs[0].text = f"Demo title {idx}"
        heading.text_frame.paragraphs[0].runs[0].font.size = Pt(32)
        content = slide.shapes.add_textbox(Inches(0.8), Inches(1.8), Inches(10), Inches(3))
        text = ("Evidence-based design improvements are achievable. " * 30) if bad else "One clear point makes a strong slide."
        content.text_frame.paragraphs[0].text = text
        content.text_frame.paragraphs[0].runs[0].font.size = Pt(10 if bad else 22)
        if bad:
            shape = slide.shapes.add_textbox(Inches(12), Inches(7.2), Inches(3), Inches(1))
            shape.text = "Off-slide text"
        if low_contrast:
            bg = slide.background.fill
            bg.solid()
            bg.fore_color.rgb = RGBColor(255, 255, 255)
            content.text_frame.paragraphs[0].runs[0].font.color.rgb = RGBColor(235, 235, 235)
    b = BytesIO()
    prs.save(b)
    return b.getvalue()


def test_clean_slide_high_score():
    result = review_pptx(make_deck(), "good.pptx")
    assert result.score >= 90
    assert result.slide_count == 1
    assert not any(f.headline == "Text under 14 pt" for f in result.slides[0].findings)


def test_bad_slide_is_flagged_and_lower_scored():
    good = review_pptx(make_deck())
    poor = review_pptx(make_deck(bad=True))
    assert poor.score < good.score
    headlines = [f.headline for f in poor.slides[0].findings]
    assert "Very dense slide" in headlines
    assert "Text under 14 pt" in headlines
    assert "Content extends outside slide" in headlines


def test_low_contrast_is_flagged_only_when_measurable():
    result = review_pptx(make_deck(low_contrast=True))
    assert result.slides[0].metrics["contrast_runs_checked"] >= 1
    assert any(f.category == "Contrast" for f in result.slides[0].findings)
    assert result.category_scores["Contrast"] < WEIGHTS["Contrast"]


def test_multiple_slides_and_json_markdown_export():
    r = review_pptx(make_deck(slides=3), "example.pptx")
    assert r.slide_count == 3
    assert len(r.slides) == 3
    assert '"score":' in r.to_json()
    assert "Slide 1:" in r.to_markdown()
    assert sum(WEIGHTS.values()) == 100
    assert r.score == round(sum(s.score for s in r.slides) / 3)


def test_contrast_math_wcag_limits():
    assert contrast_ratio((255, 255, 255), (0, 0, 0)) == 21
    assert contrast_ratio((0, 0, 0), (0, 0, 0)) == 1


@pytest.mark.parametrize("data", [b"", b"bad data", b"PK\x03\x04wrong"])
def test_rejects_invalid_upload(data):
    with pytest.raises(ValueError):
        review_pptx(data)


def test_notes_export_retains_original_and_adds_edit_feedback():
    data = make_deck(bad=True)
    report = review_pptx(data)
    annotated = add_review_notes(data, report)
    before = Presentation(BytesIO(data))
    after = Presentation(BytesIO(annotated))
    assert len(after.slides) == len(before.slides)
    assert after.slides[0].notes_slide.notes_text_frame is not None
    notes = after.slides[0].notes_slide.notes_text_frame.text
    assert "[SlideAudit]" in notes
    assert "Suggested edit:" in notes
    assert after.slides[0].shapes[0].text == before.slides[0].shapes[0].text
